#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
文档处理工具模块
支持处理 DOC、DOCX、PDF 等格式的文档
"""

import asyncio
import io
from pathlib import Path
from typing import Optional, Union

import aiofiles
from ..utils.logger import get_logger
from ..utils.decorators import monitor_performance

logger = get_logger(__name__)

# 可选依赖导入
try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    logger.warning("python-docx 未安装，无法处理 .docx 文件")

try:
    import PyPDF2
    import pdfplumber
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False
    logger.warning("PyPDF2 或 pdfplumber 未安装，无法处理 .pdf 文件")

try:
    import win32com.client
    DOC_AVAILABLE = True
    USE_WIN32 = True
except ImportError:
    DOC_AVAILABLE = False
    USE_WIN32 = False
    logger.warning("win32com 未安装，无法处理 .doc 文件（仅Windows支持）")


class DocumentProcessor:
    """文档处理器"""
    
    def __init__(self):
        self.supported_formats = []
        
        if DOCX_AVAILABLE:
            self.supported_formats.extend(['.docx'])
        if PDF_AVAILABLE:
            self.supported_formats.extend(['.pdf'])
        if DOC_AVAILABLE:
            self.supported_formats.extend(['.doc'])
        
        # 始终支持文本格式
        self.supported_formats.extend(['.txt', '.html', '.htm'])
        
        logger.info(f"文档处理器初始化完成，支持格式: {self.supported_formats}")
    
    def is_supported(self, file_path: Union[str, Path]) -> bool:
        """检查文件格式是否支持"""
        file_path = Path(file_path)
        return file_path.suffix.lower() in self.supported_formats
    
    @monitor_performance()
    async def extract_text_from_file(self, file_path: Union[str, Path]) -> Optional[str]:
        """从文件中提取文本内容"""
        file_path = Path(file_path)
        
        if not file_path.exists():
            logger.error(f"文件不存在: {file_path}")
            return None
        
        if not self.is_supported(file_path):
            logger.error(f"不支持的文件格式: {file_path.suffix}")
            return None
        
        try:
            suffix = file_path.suffix.lower()
            
            if suffix == '.docx':
                return await self._extract_from_docx(file_path)
            elif suffix == '.pdf':
                return await self._extract_from_pdf(file_path)
            elif suffix == '.doc':
                return await self._extract_from_doc(file_path)
            elif suffix in ['.txt', '.html', '.htm']:
                return await self._extract_from_text(file_path)
            else:
                logger.error(f"未实现的文件格式处理: {suffix}")
                return None
                
        except Exception as e:
            logger.error(f"提取文件 {file_path} 的文本内容失败: {e}")
            return None
    
    async def _extract_from_docx(self, file_path: Path) -> Optional[str]:
        """从DOCX文件提取文本"""
        if not DOCX_AVAILABLE:
            logger.error("python-docx 未安装，无法处理 .docx 文件")
            return None
        
        try:
            # 在线程池中执行同步操作
            loop = asyncio.get_event_loop()
            
            def _extract():
                doc = docx.Document(file_path)
                paragraphs = []
                
                # 提取段落文本
                for paragraph in doc.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        paragraphs.append(text)
                
                # 提取表格文本
                for table in doc.tables:
                    for row in table.rows:
                        row_text = []
                        for cell in row.cells:
                            cell_text = cell.text.strip()
                            if cell_text:
                                row_text.append(cell_text)
                        if row_text:
                            paragraphs.append(' | '.join(row_text))
                
                return '\n'.join(paragraphs)
            
            text = await loop.run_in_executor(None, _extract)
            logger.info(f"成功提取DOCX文件文本: {file_path.name}，长度: {len(text)}")
            return text
            
        except Exception as e:
            logger.error(f"提取DOCX文件失败 {file_path}: {e}")
            return None
    
    async def _extract_from_pdf(self, file_path: Path) -> Optional[str]:
        """从PDF文件提取文本"""
        if not PDF_AVAILABLE:
            logger.error("PyPDF2 或 pdfplumber 未安装，无法处理 .pdf 文件")
            return None
        
        try:
            # 在线程池中执行同步操作
            loop = asyncio.get_event_loop()
            
            def _extract_with_pdfplumber():
                """使用pdfplumber提取文本（推荐）"""
                text_parts = []
                with pdfplumber.open(file_path) as pdf:
                    for page_num, page in enumerate(pdf.pages, 1):
                        try:
                            page_text = page.extract_text()
                            if page_text:
                                text_parts.append(f"=== 第{page_num}页 ===\n{page_text}")
                        except Exception as e:
                            logger.warning(f"提取PDF第{page_num}页失败: {e}")
                            continue
                return '\n\n'.join(text_parts)
            
            def _extract_with_pypdf2():
                """使用PyPDF2提取文本（备用）"""
                text_parts = []
                with open(file_path, 'rb') as file:
                    pdf_reader = PyPDF2.PdfReader(file)
                    for page_num, page in enumerate(pdf_reader.pages, 1):
                        try:
                            page_text = page.extract_text()
                            if page_text:
                                text_parts.append(f"=== 第{page_num}页 ===\n{page_text}")
                        except Exception as e:
                            logger.warning(f"提取PDF第{page_num}页失败: {e}")
                            continue
                return '\n\n'.join(text_parts)
            
            # 优先使用pdfplumber
            try:
                text = await loop.run_in_executor(None, _extract_with_pdfplumber)
            except Exception as e:
                logger.warning(f"pdfplumber提取失败，尝试PyPDF2: {e}")
                text = await loop.run_in_executor(None, _extract_with_pypdf2)
            
            if text:
                logger.info(f"成功提取PDF文件文本: {file_path.name}，长度: {len(text)}")
                return text
            else:
                logger.warning(f"PDF文件无法提取文本: {file_path.name}")
                return None
                
        except Exception as e:
            logger.error(f"提取PDF文件失败 {file_path}: {e}")
            return None
    
    async def _extract_from_doc(self, file_path: Path) -> Optional[str]:
        """从DOC文件提取文本"""
        if not DOC_AVAILABLE:
            logger.error("win32com 未安装，无法处理 .doc 文件（仅Windows支持）")
            return None
        
        try:
            # 在线程池中执行同步操作
            loop = asyncio.get_event_loop()
            
            def _extract():
                # 使用win32com（仅Windows）
                import win32com.client
                word = win32com.client.Dispatch("Word.Application")
                word.Visible = False
                try:
                    doc = word.Documents.Open(str(file_path.absolute()))
                    text = doc.Content.Text
                    doc.Close()
                    return text
                finally:
                    word.Quit()
            
            text = await loop.run_in_executor(None, _extract)
            
            if text:
                logger.info(f"成功提取DOC文件文本: {file_path.name}，长度: {len(text)}")
                return text.strip()
            else:
                logger.warning(f"DOC文件无法提取文本: {file_path.name}")
                return None
                
        except Exception as e:
            logger.error(f"提取DOC文件失败 {file_path}: {e}")
            return None
    
    async def _extract_from_text(self, file_path: Path) -> Optional[str]:
        """从文本文件提取内容"""
        try:
            # 尝试多种编码
            encodings = ['utf-8', 'gbk', 'gb2312', 'utf-16', 'latin1']
            
            for encoding in encodings:
                try:
                    async with aiofiles.open(file_path, 'r', encoding=encoding) as f:
                        text = await f.read()
                        logger.info(f"成功读取文本文件: {file_path.name}，编码: {encoding}，长度: {len(text)}")
                        return text
                except UnicodeDecodeError:
                    continue
                except Exception as e:
                    logger.warning(f"使用编码 {encoding} 读取文件失败: {e}")
                    continue
            
            logger.error(f"无法使用任何编码读取文件: {file_path}")
            return None
            
        except Exception as e:
            logger.error(f"读取文本文件失败 {file_path}: {e}")
            return None
    
    async def batch_extract_text(self, file_paths: list[Union[str, Path]]) -> dict[str, Optional[str]]:
        """批量提取文本"""
        results = {}
        
        # 创建信号量限制并发数
        semaphore = asyncio.Semaphore(5)
        
        async def extract_with_semaphore(file_path):
            async with semaphore:
                return await self.extract_text_from_file(file_path)
        
        # 并发处理
        tasks = [extract_with_semaphore(fp) for fp in file_paths]
        texts = await asyncio.gather(*tasks, return_exceptions=True)
        
        # 整理结果
        for file_path, text in zip(file_paths, texts):
            if isinstance(text, Exception):
                logger.error(f"批量提取文件 {file_path} 失败: {text}")
                results[str(file_path)] = None
            else:
                results[str(file_path)] = text
        
        return results


# 全局实例
_document_processor: Optional[DocumentProcessor] = None


def get_document_processor() -> DocumentProcessor:
    """获取全局文档处理器实例"""
    global _document_processor
    if _document_processor is None:
        _document_processor = DocumentProcessor()
    return _document_processor
